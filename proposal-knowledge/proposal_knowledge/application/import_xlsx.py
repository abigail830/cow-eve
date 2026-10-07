from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from proposal_knowledge.application.jurisdiction_defaults import apply_bu_default_jurisdictions
from proposal_knowledge.application.brand_sanitize import (
    sanitize_business_unit,
    sanitize_free_text,
    sanitize_jurisdictions,
    sanitize_linked_skus,
)
from proposal_knowledge.infrastructure.db.models import PackageRow, ProductRow


def parse_jurisdictions(raw: object, *, sanitize_acorp: bool = False) -> list[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text:
        return []
    parts = re.split(r"[;,|/]+", text)
    values = [p.strip() for p in parts if p.strip()]
    if sanitize_acorp:
        return sanitize_jurisdictions(values)
    return values


def parse_linked_skus(raw: object, *, sanitize_acorp: bool = False) -> list[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text:
        return []
    values = [p.strip() for p in text.split(";") if p.strip()]
    if sanitize_acorp:
        return sanitize_linked_skus(values)
    return values


def _maybe_sanitize(text: str | None, sanitize_acorp: bool) -> str | None:
    if not sanitize_acorp or text is None:
        return text
    return sanitize_free_text(text)


def import_products_xlsx(
    session: Session, path: Path, *, sanitize_acorp: bool = False
) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    count = 0
    for row in rows:
        if not row or not any(row):
            continue
        sku = _cell(row, idx, "SKU*")
        bu = _cell(row, idx, "BU*")
        if not sku or not bu:
            continue
        if sanitize_acorp:
            bu = sanitize_business_unit(bu)
            sku = sanitize_free_text(sku) or sku
        product_name = (
            _maybe_sanitize(_cell(row, idx, "Product name*") or sku, sanitize_acorp)
            or sku
        )
        product = ProductRow(
            business_unit=bu,
            sku=sku,
            product_name=product_name,
            product_description=_maybe_sanitize(
                _cell(row, idx, "Product description"), sanitize_acorp
            ),
            service_name_on_proposal=_maybe_sanitize(
                _cell(row, idx, "Service name on Proposal"), sanitize_acorp
            ),
            scope_of_work=_maybe_sanitize(_cell(row, idx, "Scope of Work"), sanitize_acorp),
            sku_semantic_for_ai=_maybe_sanitize(
                _cell(row, idx, "SKU Semantic for AI"), sanitize_acorp
            ),
            billing_frequency=_maybe_sanitize(
                _cell(row, idx, "Billing frequency*"), sanitize_acorp
            ),
            currency=_maybe_sanitize(_cell(row, idx, "Currency*"), sanitize_acorp),
            price=_maybe_sanitize(_cell(row, idx, "Price"), sanitize_acorp),
            recurring=_maybe_sanitize(_cell(row, idx, "Recurring"), sanitize_acorp),
            standard_pricing_matrix=_maybe_sanitize(
                _cell(row, idx, "Standard pricing matrix"), sanitize_acorp
            ),
            department_team=_maybe_sanitize(
                _cell(row, idx, "Department/Team"), sanitize_acorp
            ),
            status=_maybe_sanitize(_cell(row, idx, "Status*"), sanitize_acorp),
            jurisdictions=apply_bu_default_jurisdictions(
                bu,
                parse_jurisdictions(
                    _cell(row, idx, "Jurisdictions"), sanitize_acorp=sanitize_acorp
                ),
                sku=sku,
                product_name=product_name,
            ),
        )
        session.merge(product)
        count += 1
    wb.close()
    return count


def import_packages_xlsx(
    session: Session, path: Path, *, sanitize_acorp: bool = False
) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    count = 0
    for row in rows:
        if not row or not any(row):
            continue
        package_id = _cell(row, idx, "ID*")
        bu = _cell(row, idx, "BU*")
        if not package_id or not bu:
            continue
        if sanitize_acorp:
            bu = sanitize_business_unit(bu)
            package_id = sanitize_free_text(package_id) or package_id
        pkg = PackageRow(
            business_unit=bu,
            package_id=package_id,
            package_name=_maybe_sanitize(
                _cell(row, idx, "Package name*") or package_id, sanitize_acorp
            )
            or package_id,
            package_description=_maybe_sanitize(
                _cell(row, idx, "Package description"), sanitize_acorp
            ),
            package_semantic_for_ai=_maybe_sanitize(
                _cell(row, idx, "Package Semantic for AI"), sanitize_acorp
            ),
            linked_skus=parse_linked_skus(
                _cell(row, idx, "Linked SKU"), sanitize_acorp=sanitize_acorp
            ),
        )
        session.merge(pkg)
        count += 1
    wb.close()
    return count


def _cell(row: tuple, idx: dict[str, int], col: str) -> str | None:
    i = idx.get(col)
    if i is None or i >= len(row):
        return None
    val = row[i]
    if val is None:
        return None
    text = str(val).strip()
    return text or None
