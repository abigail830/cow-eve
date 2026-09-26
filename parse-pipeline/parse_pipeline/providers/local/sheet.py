from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import load_workbook


def _csv_to_markdown(data: bytes, max_rows: int) -> tuple[str, list[str]]:
    warnings: list[str] = []
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if len(rows) > max_rows:
        warnings.append(f"truncated to {max_rows} rows")
        rows = rows[:max_rows]
    if not rows:
        return "[empty sheet]", warnings
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    for row in body:
        padded = row + [""] * (len(header) - len(row))
        lines.append("| " + " | ".join(padded[: len(header)]) + " |")
    return "\n".join(lines), warnings


def _xlsx_to_markdown(data: bytes, max_rows: int) -> tuple[str, list[str]]:
    warnings: list[str] = []
    wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[list[Any]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                warnings.append(f"sheet {sheet_name!r} truncated to {max_rows} rows")
                break
            rows.append(["" if cell is None else str(cell) for cell in row])
        if not rows:
            continue
        parts.append(f"## {sheet_name}\n")
        header = rows[0]
        body = rows[1:] if len(rows) > 1 else []
        col_count = max(len(header), 1)
        parts.append("| " + " | ".join(header) + " |")
        parts.append("| " + " | ".join(["---"] * col_count) + " |")
        for row in body:
            padded = row + [""] * (col_count - len(row))
            parts.append("| " + " | ".join(padded[:col_count]) + " |")
        parts.append("")
    wb.close()
    content = "\n".join(parts).strip()
    if not content:
        return "[empty workbook]", warnings
    return content, warnings


def extract_sheet_bytes(data: bytes, filename: str | None, max_rows: int) -> tuple[str, list[str]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _csv_to_markdown(data, max_rows)
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        if name.endswith(".xls"):
            raise ValueError("legacy .xls not supported in local sheet parser; use document_mind")
        return _xlsx_to_markdown(data, max_rows)
    # guess csv
    if b"," in data[:4096] and b"\n" in data[:4096]:
        return _csv_to_markdown(data, max_rows)
    return _xlsx_to_markdown(data, max_rows)
